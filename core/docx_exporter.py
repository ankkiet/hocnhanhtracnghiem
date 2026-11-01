import base64
import html
import io
import os
import re
from typing import Optional

import latex2mathml.converter
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

# Khởi tạo XSLT Transformer từ file MML2OMML.XSL chuẩn của Microsoft
XSL_PATH = os.path.join(os.path.dirname(__file__), 'MML2OMML.XSL')
XSLT_TRANSFORMER = None
try:
    if os.path.exists(XSL_PATH):
        xslt_doc = etree.parse(XSL_PATH)
        XSLT_TRANSFORMER = etree.XSLT(xslt_doc)
    else:
        print(f"Warning: MML2OMML.XSL not found at {XSL_PATH}")
except Exception as e:
    print(f"Warning: Failed to load MML2OMML.XSL: {e}")


# Danh sách các lệnh LaTeX chuẩn bắt đầu bằng chữ 'n' không được làm biến dạng
VALID_LATEX_N_COMMANDS = {
    'nearrow', 'nwarrow', 'neq', 'ne', 'notin', 'nabla', 'nu', 'not', 'neg',
    'natural', 'nleq', 'ngeq', 'nless', 'ngtr', 'nprec', 'nsucc', 'nsim',
    'nsubset', 'nsupset', 'nsubseteq', 'nsupseteq', 'normalsize', 'noindent',
    'newcommand', 'newenvironment'
}


def sanitize_latex_formula(latex_str: str) -> str:
    """
    Làm sạch và phục hồi công thức LaTeX bị lỗi HTML/newline/bảng biến thiên trước khi render sang OMML/MathML.
    Khắc phục triệt để lỗi thẻ <br>, chuỗi thoát \\n, dấu \\ thừa làm vỡ bảng biến thiên và công thức.
    """
    if not latex_str or not latex_str.strip():
        return ""
    s = latex_str.strip()

    # 1. Gỡ bỏ các ký hiệu bao quanh ($$, $, \[, \])
    if s.startswith('$$') and s.endswith('$$') and len(s) >= 4:
        s = s[2:-2].strip()
    elif s.startswith('$') and s.endswith('$') and len(s) >= 2:
        s = s[1:-1].strip()
    elif s.startswith(r'\[') and s.endswith(r'\]') and len(s) >= 4:
        s = s[2:-2].strip()
    elif s.startswith(r'\(') and s.endswith(r'\)') and len(s) >= 4:
        s = s[2:-2].strip()

    # 2. Chuyển đổi các thực thể HTML phổ biến
    s = s.replace('&nbsp;', ' ').replace('&lt;', '<').replace('&gt;', '>').replace('&amp;', '&')

    # 3. Xử lý các thẻ <br>, <br/> bị lọt vào công thức
    # Nếu đi liền sau dấu \\ (xuống dòng của LaTeX) -> Giữ lại \\ và loại bỏ <br>
    s = re.sub(r'\\\\(?:\s*<br\s*\/?>|\s*\\?n|\s*\\?\r?\n)+', r'\\\\ \n', s, flags=re.IGNORECASE)
    s = re.sub(r'\\<br\s*\/?>', r'\\\\ \n', s, flags=re.IGNORECASE)

    # Nếu là <br> đứng trong môi trường bảng/mảng (array, matrix, cases, aligned)
    if any(env in s for env in ['array', 'matrix', 'pmatrix', 'bmatrix', 'cases', 'align', 'aligned']):
        # Nếu <br> ở ngay trước \end{...}, loại bỏ để tránh tạo dòng trống cuối bảng
        s = re.sub(r'<br\s*\/?>\s*(?:\\?n\s*)*(?=\\end\{)', ' ', s, flags=re.IGNORECASE)
        # Các <br> còn lại trong mảng chuyển thành \\ xuống dòng LaTeX
        s = re.sub(r'<br\s*\/?>', r'\\\\ \n', s, flags=re.IGNORECASE)
    else:
        s = re.sub(r'<br\s*\/?>', ' ', s, flags=re.IGNORECASE)

    # Loại bỏ các thẻ HTML bao bọc khác nếu có
    s = re.sub(r'<\/?(?:span|div|p|b|i|u|font|mark)[^>]*>', ' ', s, flags=re.IGNORECASE)

    # 4. Khắc phục lỗi chuỗi thoát \\n hoặc \n (bị in ra thành text \nx, \ny, \n-1...)
    def _fix_n_command(m):
        cmd = m.group(1)
        if cmd.lower() in VALID_LATEX_N_COMMANDS:
            return '\\' + cmd
        if cmd.startswith('n') and len(cmd) > 1:
            return ' ' + cmd[1:]
        if cmd == 'n':
            return ' '
        return '\\' + cmd

    s = re.sub(r'\\([a-zA-Z]+)', _fix_n_command, s)
    s = re.sub(r'\\n(?=[^a-zA-Z]|$)', ' ', s)

    # 5. Dọn dẹp dấu \\ thừa ở cuối môi trường (tránh sinh dòng trống cuối bảng)
    s = re.sub(r'\\\\\s*\\end\{(array|matrix|pmatrix|bmatrix|cases|aligned)\}', r'\\end{\1}', s)

    # 6. Thay thế macro đặc thù
    s = s.replace(r'\degree', r'^\circ')

    return s.strip()


def latex_to_clean_unicode(latex_str: str) -> str:
    """Chuyển đổi chuỗi LaTeX sang Unicode toán học dễ đọc (dùng khi fallback)."""
    if not latex_str:
        return ""
    s = sanitize_latex_formula(latex_str)
    if not s:
        return ""

    # Phân số \frac{a}{b} -> (a)/(b)
    for _ in range(4):
        if r'\frac' not in s:
            break
        s = re.sub(r'\\frac\{([^{}]+)\}\{([^{}]+)\}', r'(\1)/(\2)', s)

    # Căn thức
    s = re.sub(r'\\sqrt\[([^\]]+)\]\{([^{}]+)\}', r'[\1]√(\2)', s)
    s = re.sub(r'\\sqrt\{([^{}]+)\}', r'√(\1)', s)

    # Chữ cái Hy Lạp
    greek_map = {
        r'\alpha': 'α', r'\beta': 'β', r'\gamma': 'γ', r'\Delta': 'Δ', r'\delta': 'δ',
        r'\epsilon': 'ε', r'\varepsilon': 'ε', r'\zeta': 'ζ', r'\eta': 'η', r'\theta': 'θ',
        r'\lambda': 'λ', r'\mu': 'μ', r'\pi': 'π', r'\rho': 'ρ', r'\sigma': 'σ',
        r'\Sigma': 'Σ', r'\tau': 'τ', r'\phi': 'φ', r'\varphi': 'ϕ', r'\omega': 'ω',
        r'\Omega': 'Ω'
    }
    for k, v in greek_map.items():
        s = re.sub(re.escape(k) + r'(?![a-zA-Z])', v, s)

    # Ký hiệu toán học
    sym_map = {
        r'\pm': '±', r'\mp': '∓', r'\times': '×', r'\div': '÷', r'\cdot': '·',
        r'\leq': '≤', r'\le': '≤', r'\geq': '≥', r'\ge': '≥', r'\neq': '≠', r'\ne': '≠',
        r'\approx': '≈', r'\equiv': '≡', r'\in': '∈', r'\notin': '∉', r'\subset': '⊂',
        r'\cup': '∪', r'\cap': '∩', r'\emptyset': '∅', r'\infty': '∞',
        r'\rightarrow': '→', r'\to': '→', r'\longrightarrow': '⟶',
        r'\Rightarrow': '⇒', r'\Leftrightarrow': '⇔', r'\rightleftharpoons': '⇌',
        r'\forall': '∀', r'\exists': '∃', r'\perp': '⊥', r'\parallel': '∥',
        r'\angle': '∠', r'\circ': '°', r'\degree': '°',
        r'\dots': '...', r'\ldots': '...', r'\cdots': '...',
        r'\setminus': '\\', r'\mathbb{R}': 'ℝ', r'\mathbb{N}': 'ℕ',
        r'\mathbb{Z}': 'ℤ', r'\mathbb{Q}': 'ℚ', r'\mathbb{C}': 'ℂ',
        r'\int': '∫', r'\sum': '∑'
    }
    for k, v in sym_map.items():
        s = re.sub(re.escape(k) + r'(?![a-zA-Z])', v, s)

    # Bỏ các lệnh bọc text/font
    s = re.sub(r'\\(?:text|mathrm|mathbf|mathit|textbf|textit|underline|overline)\{([^{}]*)\}', r'\1', s)
    s = re.sub(r'\\left([(\[{|])', r'\1', s)
    s = re.sub(r'\\right([)\]}|])', r'\1', s)
    s = s.replace(r'\{', '{').replace(r'\}', '}')
    s = s.replace(r'\,', ' ').replace(r'\;', ' ').replace(r'\:', ' ').replace(r'\!', '')

    # Chỉ số trên và dưới đơn giản
    sup_map = {'0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴', '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹', '+': '⁺', '-': '⁻', 'n': 'ⁿ', 'x': 'ˣ'}
    sub_map = {'0': '₀', '1': '₁', '2': '₂', '3': '₃', '4': '₄', '5': '₅', '6': '₆', '7': '₇', '8': '₈', '9': '₉', '+': '₊', '-': '₋', 'x': 'ₓ'}
    for k, v in sup_map.items():
        s = re.sub(r'\^\{?' + re.escape(k) + r'\}?', v, s)
    for k, v in sub_map.items():
        s = re.sub(r'_\{?' + re.escape(k) + r'\}?', v, s)

    s = re.sub(r'[{}]', '', s)
    s = s.replace('\\\\', ' ')
    return s.strip()


def latex_to_omml(latex_str: str) -> Optional[OxmlElement]:
    """Chuyển đổi chuỗi công thức LaTeX sang thẻ <m:oMath> chuẩn của Microsoft Word."""
    if not XSLT_TRANSFORMER or not latex_str or not latex_str.strip():
        return None

    s = sanitize_latex_formula(latex_str)
    if not s:
        return None

    try:

        # 1. Chuyển LaTeX sang MathML
        mathml = latex2mathml.converter.convert(s)
        # Sửa các ký tự '&' chưa escape trong MathML (như aligned hoặc ma trận)
        mathml_fixed = re.sub(r'&(?!(?:amp|lt|gt|quot|apos|#x?[0-9a-fA-F]+);)', '&amp;', mathml)

        # 2. Parse MathML thành XML DOM với parser có tính năng tự phục hồi
        parser = etree.XMLParser(recover=True)
        dom = etree.fromstring(mathml_fixed.encode('utf-8'), parser=parser)

        # 3. Dùng XSLT biến đổi MathML sang Office MathML (OMML)
        new_dom = XSLT_TRANSFORMER(dom)
        omml_bytes = etree.tostring(new_dom, encoding='utf-8')

        # 4. Parse thành OxmlElement của docx
        elem = parse_xml(omml_bytes)

        # Nếu phần tử bọc ngoài là oMathPara, lấy oMath bên trong để hiển thị inline liền mạch
        if elem.tag.endswith('oMathPara'):
            inner = elem.find(qn('m:oMath'))
            if inner is not None:
                return inner
        return elem
    except Exception:
        return None


DOCX_IMAGE_CACHE = {}


def fetch_image_stream(src: str) -> Optional[io.BytesIO]:
    """Tải hoặc giải mã dữ liệu ảnh từ URL, Base64, hoặc file path."""
    try:
        if not src:
            return None
        if src.startswith("data:image/"):
            header, encoded = src.split(",", 1)
            img_bytes = base64.b64decode(encoded)
            return io.BytesIO(img_bytes)
        elif src.startswith("http://") or src.startswith("https://"):
            if src in DOCX_IMAGE_CACHE:
                return io.BytesIO(DOCX_IMAGE_CACHE[src])
            import requests
            resp = requests.get(src, timeout=5)
            if resp.status_code == 200:
                DOCX_IMAGE_CACHE[src] = resp.content
                return io.BytesIO(resp.content)
        elif os.path.exists(src):
            with open(src, "rb") as f:
                return io.BytesIO(f.read())
    except Exception as e:
        print(f"Warning: Không thể tải ảnh '{src[:50]}...': {e}")
    return None


def add_formatted_text(
    paragraph,
    text: str,
    font_name: str = "Times New Roman",
    font_size_pt: float = 11,
    base_bold: bool = False,
    base_italic: bool = False,
    base_color: Optional[RGBColor] = None
):
    """
    Phân tích chuỗi chứa text, công thức LaTeX, thẻ HTML (sub, sup, b, i, img, br)
    và xuất ra paragraph trong Word dưới dạng công thức toán học chuyên nghiệp (OMML)
    và định dạng chuẩn in ấn.
    """
    if not text:
        return

    # Pattern nhận diện:
    # 1. Khối công thức Toán LaTeX: $$, \[, \(, $...$
    # 2. Khối LaTeX trần thường gặp: \frac{...}{...}, \sqrt{...}
    # 3. Thẻ ảnh HTML: <img ...>
    # 4. Xuống dòng: <br>, \n
    # 5. Thẻ định dạng HTML: <b>, </b>, <i>, </i>, <u>, </u>, <sup>, </sup>, <sub>, </sub>, <p>, </p>
    pattern = re.compile(
        r'('
        r'\$\$.*?\$\$|'
        r'\\\[.*?\\\]|'
        r'\\\(.*?\\\)|'
        r'\\begin\{(?:array|matrix|pmatrix|bmatrix|cases|aligned|align\*?)\}[\s\S]*?\\end\{(?:array|matrix|pmatrix|bmatrix|cases|aligned|align\*?)\}|'
        r'(?<!\\)\$(?!\s)[^\$\n]+?(?<!\s)(?<!\\)\$|'
        r'\\(?:frac|sqrt)\s*(?:\[[^\]]*\]\s*)?\{[^{}]*\}\s*(?:\{[^{}]*\})?|'
        r'<img\b[^>]*\/?>|'
        r'<br\s*\/?>|'
        r'</?[a-zA-Z][^>]*>|'
        r'\r?\n'
        r')',
        re.DOTALL | re.IGNORECASE
    )

    parts = pattern.split(text)

    # Trạng thái định dạng
    is_bold = base_bold
    is_italic = base_italic
    is_underline = False
    is_sup = False
    is_sub = False

    for part in parts:
        if not part:
            continue

        lower_part = part.lower().strip()

        # 1. Thẻ HTML xuống dòng hoặc ký tự newline
        if lower_part in ('<br>', '<br/>', '<br />', '\n', '\r\n') or part in ('\n', '\r\n'):
            run = paragraph.add_run()
            run.add_break()
            continue

        # 2. Xử lý các thẻ định dạng HTML
        if lower_part in ('<b>', '<strong>'):
            is_bold = True
            continue
        elif lower_part in ('</b>', '</strong>'):
            is_bold = base_bold
            continue
        elif lower_part in ('<i>', '<em>'):
            is_italic = True
            continue
        elif lower_part in ('</i>', '</em>'):
            is_italic = base_italic
            continue
        elif lower_part == '<u>':
            is_underline = True
            continue
        elif lower_part == '</u>':
            is_underline = False
            continue
        elif lower_part == '<sup>':
            is_sup = True
            continue
        elif lower_part == '</sup>':
            is_sup = False
            continue
        elif lower_part == '<sub>':
            is_sub = True
            continue
        elif lower_part == '</sub>':
            is_sub = False
            continue
        elif lower_part in ('<p>', '</p>', '<span>', '</span>'):
            continue

        # 3. Thẻ ảnh HTML <img src="...">
        if lower_part.startswith('<img'):
            src_match = re.search(r'src=["\']([^"\']+)["\']', part, re.IGNORECASE)
            if src_match:
                src = src_match.group(1)
                stream = fetch_image_stream(src)
                if stream:
                    try:
                        run_img = paragraph.add_run()
                        run_img.add_picture(stream, width=Inches(3.5))
                        continue
                    except Exception as e:
                        print(f"Error adding picture to docx: {e}")
            continue

        # 4. Công thức toán học LaTeX
        is_math = (
            part.startswith('$$') and part.endswith('$$') and len(part) >= 4 or
            part.startswith(r'\[') and part.endswith(r'\]') and len(part) >= 4 or
            part.startswith(r'\(') and part.endswith(r'\)') and len(part) >= 4 or
            part.startswith('$') and part.endswith('$') and len(part) >= 2 or
            part.startswith(r'\frac') or
            part.startswith(r'\sqrt') or
            part.startswith(r'\begin{')
        )

        if is_math:
            omml_elem = latex_to_omml(part)
            if omml_elem is not None:
                # Chèn trực tiếp Equation Office MathML vào đoạn văn
                paragraph._p.append(omml_elem)
            else:
                # Fallback: Chuyển sang ký tự Unicode toán học sạch sẽ
                clean_math = latex_to_clean_unicode(part)
                run = paragraph.add_run(clean_math)
                run.font.name = font_name
                run.font.size = Pt(font_size_pt)
                run.italic = True
                if base_color:
                    run.font.color.rgb = base_color
            continue

        # 5. Đoạn text thông thường
        unescaped_text = html.unescape(part)
        # Loại bỏ dấu backslash thừa ở cuối dòng text (ví dụ: "có bảng biến thiên như sau:\")
        unescaped_text = re.sub(r'\\+(?=\s*$)', '', unescaped_text)
        run = paragraph.add_run(unescaped_text)
        run.font.name = font_name
        run.font.size = Pt(font_size_pt)
        run.bold = is_bold
        run.italic = is_italic
        run.underline = is_underline
        if is_sup:
            run.font.superscript = True
        if is_sub:
            run.font.subscript = True
        if base_color:
            run.font.color.rgb = base_color


def render_html_table_in_docx(doc, table_html: str, font_name: str = "Times New Roman", font_size_pt: float = 10.5):
    """Chuyển đổi thẻ HTML <table> thành bảng Word chuyên nghiệp."""
    try:
        from lxml import html as lxml_html
        tree = lxml_html.fragment_fromstring(table_html, create_parent='div')
        table_nodes = tree.xpath('.//table')
        if not table_nodes:
            return
        tbl = table_nodes[0]
        rows = tbl.xpath('.//tr')
        if not rows:
            return

        grid_data = []
        max_cols = 0
        for tr in rows:
            cells = tr.xpath('./td | ./th')
            cell_texts = []
            for td in cells:
                raw_cell = lxml_html.tostring(td, encoding='utf-8').decode('utf-8')
                inner = re.sub(r'^<t[dh][^>]*>', '', raw_cell, flags=re.IGNORECASE)
                inner = re.sub(r'</t[dh]>$', '', inner, flags=re.IGNORECASE).strip()
                cell_texts.append(inner)
            if cell_texts:
                grid_data.append(cell_texts)
                if len(cell_texts) > max_cols:
                    max_cols = len(cell_texts)

        if not grid_data or max_cols == 0:
            return

        word_tbl = doc.add_table(rows=len(grid_data), cols=max_cols)
        word_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        word_tbl.style = 'Table Grid'

        for r_idx, row_cells in enumerate(grid_data):
            trPr = word_tbl.rows[r_idx]._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for c_idx, cell_content in enumerate(row_cells):
                cell = word_tbl.cell(r_idx, c_idx)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.05
                add_formatted_text(p, cell_content, font_name=font_name, font_size_pt=font_size_pt)
    except Exception as e:
        print(f"Warning: Không thể render HTML table sang docx: {e}")


def add_block_content(doc, text: str, font_name: str = "Times New Roman", font_size_pt: float = 11,
                      first_p=None, space_before: Pt = Pt(2), space_after: Pt = Pt(2), line_spacing: float = 1.05):
    """
    Hỗ trợ xuất nội dung có thể chứa cả thẻ HTML <table> lẫn text/LaTeX vào tài liệu Word.
    Nếu có thẻ <table>, tự động render thành bảng Word căn giữa chuyên nghiệp.
    """
    if not text:
        return
    if "<table" not in text.lower():
        p = first_p if first_p is not None else doc.add_paragraph()
        if first_p is None:
            p.paragraph_format.space_before = space_before
            p.paragraph_format.space_after = space_after
            p.paragraph_format.line_spacing = line_spacing
        add_formatted_text(p, text, font_name=font_name, font_size_pt=font_size_pt)
        return

    table_pattern = re.compile(r'(<table\b[\s\S]*?<\/table>)', re.IGNORECASE)
    parts = table_pattern.split(text)
    is_first = True
    for part in parts:
        if not part or not part.strip():
            continue
        if part.lower().startswith('<table'):
            render_html_table_in_docx(doc, part, font_name=font_name, font_size_pt=font_size_pt - 0.5)
        else:
            if is_first and first_p is not None:
                p = first_p
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = line_spacing
            add_formatted_text(p, part, font_name=font_name, font_size_pt=font_size_pt)
            is_first = False


def create_borderless_table(doc, num_rows, num_cols, col_widths=None):
    tbl = doc.add_table(rows=num_rows, cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = tbl._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/></w:tblBorders>'
    )
    tblPr.append(tblBorders)
    tblCellMar = parse_xml(
        f'<w:tblCellMar {nsdecls("w")}><w:top w:w="0" w:type="dxa"/><w:bottom w:w="30" w:type="dxa"/><w:left w:w="40" w:type="dxa"/><w:right w:w="40" w:type="dxa"/></w:tblCellMar>'
    )
    tblPr.append(tblCellMar)
    for row in tbl.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                if idx < len(row.cells):
                    row.cells[idx].width = width
    return tbl


def determine_options_columns(options: list, user_choice: str = "auto") -> int:
    if not options:
        return 1
    if user_choice in ("1", "1_col", 1):
        return 1
    if user_choice in ("4", "4_col", 4) and len(options) == 4:
        return 4
    if user_choice in ("2", "2_col", 2) and len(options) in (2, 4):
        return 2

    for opt in options:
        opt_s = str(opt)
        if '<img' in opt_s.lower() or '$$' in opt_s or r'\[' in opt_s or '\n' in opt_s:
            return 1

    def clean_len(opt):
        s = str(opt)
        s = re.sub(r'<[^>]+>', '', s)
        s = re.sub(r'\\(?:frac|sqrt|left|right|text|mathrm)\b', '', s)
        s = re.sub(r'[\$\{\}\\\^_\s]+', ' ', s)
        return len(s.strip())

    lengths = [clean_len(opt) for opt in options]
    max_len = max(lengths) if lengths else 0

    if len(options) == 4:
        if max_len <= 18:
            return 4
        elif max_len <= 42:
            return 2
        else:
            return 1
    elif len(options) == 2:
        if max_len <= 42:
            return 2
        else:
            return 1
    return 1


def is_option_correct(opt_idx: int, opt_text: str, raw_ca) -> bool:
    if not raw_ca:
        return False
    clean_opt = str(opt_text).strip()
    if clean_opt.startswith('*'):
        return True
    clean_ca = str(raw_ca).strip()
    letter = chr(ord('A') + opt_idx) if 0 <= opt_idx < 26 else ""
    if clean_ca.upper() == letter:
        return True
    if clean_ca.upper().startswith(f"{letter}.") or clean_ca.upper().startswith(f"{letter}:") or clean_ca.upper().startswith(f"{letter})"):
        return True
    if clean_ca == str(opt_idx):
        return True
    match_prefix = re.match(r'^(\*?\s*[A-Fa-f][\.\:\)])\s*(.*)', clean_opt)
    if match_prefix and match_prefix.group(2).strip() and clean_ca.strip() == match_prefix.group(2).strip():
        return True
    return False


def export_quiz_to_docx(
    title: str,
    questions: Optional[list] = None,
    time_limit: int = 0,
    target_audience: str = "teacher",
    layout_mode: str = "compact",
    options_layout: str = "auto",
    show_answers: Optional[bool] = None,
    show_key_table: Optional[bool] = None,
    show_explains: Optional[bool] = None,
    include_student_answer_grid: bool = False
) -> io.BytesIO:
    if isinstance(title, dict):
        quiz_data = title
        title = quiz_data.get('title', 'Đề thi trắc nghiệm')
        questions = quiz_data.get('data') or quiz_data.get('questions') or []
        time_limit = quiz_data.get('time_limit', 0)
    elif questions is None:
        questions = []

    doc = Document()
    is_compact = (layout_mode == "compact")

    # 1. Cấu hình lề trang (Margins)
    for section in doc.sections:
        if is_compact:
            # Lề tiết kiệm giấy (Narrow Margins ~ 1.27cm - 1.4cm)
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.55)
            section.right_margin = Inches(0.55)
        else:
            # Lề tiêu chuẩn (Standard Margins ~ 1.9cm - 2.0cm)
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

    # 2. Xử lý các cờ hiển thị đáp án
    if show_answers is None:
        show_answers = (target_audience != "student")
    if show_key_table is None:
        show_key_table = (target_audience != "student")
    if show_explains is None:
        show_explains = (target_audience != "student")

    # Kích thước chữ & khoảng cách
    font_size_pt = 10.5 if is_compact else 11.0
    q_space_before = Pt(3 if is_compact else 6)
    q_space_after = Pt(1.5 if is_compact else 3)
    q_line_spacing = 1.08 if is_compact else 1.2
    opt_space_after = Pt(1.5 if is_compact else 2.5)
    opt_line_spacing = 1.05 if is_compact else 1.15
    available_width_inches = 7.17 if is_compact else 6.77

    # 3. Tiêu đề đề thi
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header.paragraph_format.space_before = Pt(0)
    p_header.paragraph_format.space_after = Pt(3 if is_compact else 4)
    run_header = p_header.add_run(f"ĐỀ KIỂM TRA TRẮC NGHIỆM: {title.upper()}")
    run_header.bold = True
    run_header.font.name = 'Times New Roman'
    run_header.font.size = Pt(13 if is_compact else 14)
    run_header.font.color.rgb = RGBColor(26, 86, 219)

    if target_audience == "teacher":
        p_badge = doc.add_paragraph()
        p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_badge.paragraph_format.space_before = Pt(0)
        p_badge.paragraph_format.space_after = Pt(3)
        run_badge = p_badge.add_run("[ BẢN GIÁO VIÊN - ĐÁP ÁN & HƯỚNG DẪN GIẢI ]")
        run_badge.bold = True
        run_badge.font.name = 'Times New Roman'
        run_badge.font.size = Pt(10.5)
        run_badge.font.color.rgb = RGBColor(185, 28, 28)

    # Thông tin thời gian và học sinh
    time_str = f"Thời gian làm bài: {time_limit} phút" if time_limit > 0 else "Thời gian làm bài: Tự do"
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_info.paragraph_format.space_before = Pt(0)
    p_info.paragraph_format.space_after = Pt(3 if is_compact else 4)
    run_time = p_info.add_run(f"{time_str} | Tổng số câu: {len(questions)} câu\n")
    run_time.font.name = 'Times New Roman'
    run_time.font.size = Pt(font_size_pt)
    run_time.italic = True

    run_stu = p_info.add_run("Họ và tên thí sinh: .............................................................. SBD: ................. Lớp: .........................")
    run_stu.font.name = 'Times New Roman'
    run_stu.font.size = Pt(font_size_pt)
    run_stu.bold = True

    # Phiếu điền đáp án trắc nghiệm cho học sinh (nếu được bật)
    if target_audience == "student" and include_student_answer_grid and questions:
        p_grid_title = doc.add_paragraph()
        p_grid_title.paragraph_format.space_before = Pt(4)
        p_grid_title.paragraph_format.space_after = Pt(2)
        run_gt = p_grid_title.add_run("PHIẾU TRẢ LỜI TRẮC NGHIỆM:")
        run_gt.bold = True
        run_gt.font.name = 'Times New Roman'
        run_gt.font.size = Pt(10)
        run_gt.font.color.rgb = RGBColor(30, 41, 59)

        total_q = len(questions)
        grid_cols = 20 if total_q > 40 else 10
        grid_rows = ((total_q + grid_cols - 1) // grid_cols) * 2
        grid_table = doc.add_table(rows=grid_rows, cols=grid_cols)
        grid_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        grid_table.style = 'Table Grid'
        for row in grid_table.rows:
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        for q_idx in range(total_q):
            row_pair = (q_idx // grid_cols) * 2
            col_pos = q_idx % grid_cols
            # Ô nhãn câu
            c_header = grid_table.cell(row_pair, col_pos)
            c_header.text = f"{q_idx + 1}"
            p_ch = c_header.paragraphs[0]
            p_ch.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_ch.paragraph_format.space_before = Pt(1)
            p_ch.paragraph_format.space_after = Pt(1)
            if p_ch.runs:
                p_ch.runs[0].bold = True
                p_ch.runs[0].font.name = 'Times New Roman'
                p_ch.runs[0].font.size = Pt(8.5 if grid_cols == 20 else 9.5)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
            c_header._tc.get_or_add_tcPr().append(shd)

            # Ô để trống cho học sinh viết
            c_empty = grid_table.cell(row_pair + 1, col_pos)
            c_empty.text = " "
            p_ce = c_empty.paragraphs[0]
            p_ce.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_ce.paragraph_format.space_before = Pt(1)
            p_ce.paragraph_format.space_after = Pt(2)

    p_div = doc.add_paragraph("-" * (65 if is_compact else 55))
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(2)
    p_div.paragraph_format.space_after = Pt(6 if is_compact else 10)

    # 4. Danh sách câu hỏi
    current_group = None
    for idx, q in enumerate(questions, 1):
        group_title = q.get('group_title', '').strip()
        if group_title and group_title != current_group:
            current_group = group_title
            p_group = doc.add_paragraph()
            p_group.paragraph_format.space_before = Pt(6 if is_compact else 8)
            p_group.paragraph_format.space_after = Pt(2 if is_compact else 4)
            run_g = p_group.add_run(f"[{group_title}]")
            run_g.bold = True
            run_g.italic = True
            run_g.font.name = 'Times New Roman'
            run_g.font.size = Pt(font_size_pt)
            run_g.font.color.rgb = RGBColor(71, 85, 105)

        q_text = q.get('question', '')
        options = q.get('options', [])
        q_type = q.get('type', 'mcq')
        raw_ca = q.get('correct_answer')

        # Dòng đề bài
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = q_space_before
        p_q.paragraph_format.space_after = q_space_after
        p_q.paragraph_format.line_spacing = q_line_spacing
        p_q.paragraph_format.keep_with_next = True

        run_q_title = p_q.add_run(f"Câu {idx}: ")
        run_q_title.bold = True
        run_q_title.font.name = 'Times New Roman'
        run_q_title.font.size = Pt(font_size_pt)
        run_q_title.font.color.rgb = RGBColor(30, 41, 59)

        add_block_content(doc, q_text, font_name='Times New Roman', font_size_pt=font_size_pt, first_p=p_q, line_spacing=q_line_spacing)

        # Xử lý các dạng câu hỏi
        if q_type == 'short_answer':
            p_ans = doc.add_paragraph()
            p_ans.paragraph_format.left_indent = Inches(0.25 if is_compact else 0.3)
            p_ans.paragraph_format.space_before = Pt(1 if is_compact else 3)
            p_ans.paragraph_format.space_after = Pt(2 if is_compact else 4)
            run_prompt = p_ans.add_run("Đáp số: ")
            run_prompt.bold = True
            run_prompt.font.name = 'Times New Roman'
            run_prompt.font.size = Pt(font_size_pt)
            if show_answers:
                run_ca = p_ans.add_run(f" {raw_ca or ''}")
                run_ca.bold = True
                run_ca.font.name = 'Times New Roman'
                run_ca.font.size = Pt(font_size_pt)
                run_ca.font.color.rgb = RGBColor(185, 28, 28)
            else:
                run_dots = p_ans.add_run("...................................................................................")
                run_dots.font.name = 'Times New Roman'
                run_dots.font.size = Pt(font_size_pt)

        elif q_type == 'true_false' or isinstance(raw_ca, dict):
            raw_ca_dict = raw_ca if isinstance(raw_ca, dict) else {}
            for opt in options:
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.left_indent = Inches(0.25 if is_compact else 0.3)
                p_opt.paragraph_format.space_before = Pt(0)
                p_opt.paragraph_format.space_after = opt_space_after
                p_opt.paragraph_format.line_spacing = opt_line_spacing

                match_prefix = re.match(r'^(\*?\s*([a-dA-D]))([\.\:\)])\s*(.*)', str(opt).strip())
                sub_key = ""
                if match_prefix:
                    sub_key = match_prefix.group(2).lower()
                    label = match_prefix.group(1) + match_prefix.group(3) + " "
                    content = match_prefix.group(4)
                    run_label = p_opt.add_run(label)
                    run_label.bold = True
                    run_label.font.name = 'Times New Roman'
                    run_label.font.size = Pt(font_size_pt)
                    add_formatted_text(p_opt, content, font_name='Times New Roman', font_size_pt=font_size_pt)
                else:
                    add_formatted_text(p_opt, str(opt), font_name='Times New Roman', font_size_pt=font_size_pt)

                if show_answers and sub_key in raw_ca_dict:
                    val = raw_ca_dict.get(sub_key)
                    is_d = (val is True or str(val).lower() in ('đ', 'd', 'true', '1'))
                    tag_text = f"   [ Đúng / Sai ] - {'ĐÚNG' if is_d else 'SAI'}"
                    run_tag = p_opt.add_run(tag_text)
                    run_tag.bold = True
                    run_tag.font.name = 'Times New Roman'
                    run_tag.font.size = Pt(font_size_pt - 1)
                    run_tag.font.color.rgb = RGBColor(22, 101, 52) if is_d else RGBColor(185, 28, 28)
                else:
                    run_tag = p_opt.add_run("   [ Đúng / Sai ]")
                    run_tag.italic = True
                    run_tag.bold = True
                    run_tag.font.name = 'Times New Roman'
                    run_tag.font.size = Pt(font_size_pt - 1)
                    run_tag.font.color.rgb = RGBColor(100, 116, 139)

        else:
            cols = determine_options_columns(options, options_layout)

            if cols in (2, 4) and len(options) in (2, 4):
                num_cols = cols
                num_rows = len(options) // num_cols
                col_width = Inches(available_width_inches / num_cols)
                table = create_borderless_table(doc, num_rows, num_cols, [col_width] * num_cols)

                for opt_idx, opt in enumerate(options):
                    r_idx = opt_idx // num_cols
                    c_idx = opt_idx % num_cols
                    cell = table.cell(r_idx, c_idx)
                    p_cell = cell.paragraphs[0]
                    p_cell.paragraph_format.space_before = Pt(0)
                    p_cell.paragraph_format.space_after = opt_space_after
                    p_cell.paragraph_format.line_spacing = opt_line_spacing

                    is_corr = is_option_correct(opt_idx, str(opt), raw_ca)
                    match_prefix = re.match(r'^(\*?\s*[A-Fa-f][\.\:\)])\s*(.*)', str(opt).strip())
                    if match_prefix:
                        label = match_prefix.group(1) + " "
                        content = match_prefix.group(2)
                        run_label = p_cell.add_run(label)
                        run_label.bold = True
                        run_label.font.name = 'Times New Roman'
                        run_label.font.size = Pt(font_size_pt)
                        if show_answers and is_corr:
                            run_label.underline = True
                            run_label.font.color.rgb = RGBColor(185, 28, 28)
                        add_formatted_text(p_cell, content, font_name='Times New Roman', font_size_pt=font_size_pt)
                    else:
                        add_formatted_text(p_cell, str(opt), font_name='Times New Roman', font_size_pt=font_size_pt)

            else:
                for opt_idx, opt in enumerate(options):
                    p_opt = doc.add_paragraph()
                    p_opt.paragraph_format.left_indent = Inches(0.25 if is_compact else 0.3)
                    p_opt.paragraph_format.space_before = Pt(0)
                    p_opt.paragraph_format.space_after = opt_space_after
                    p_opt.paragraph_format.line_spacing = opt_line_spacing

                    is_corr = is_option_correct(opt_idx, str(opt), raw_ca)
                    match_prefix = re.match(r'^(\*?\s*[A-Fa-f][\.\:\)])\s*(.*)', str(opt).strip())
                    if match_prefix:
                        label = match_prefix.group(1) + " "
                        content = match_prefix.group(2)
                        run_label = p_opt.add_run(label)
                        run_label.bold = True
                        run_label.font.name = 'Times New Roman'
                        run_label.font.size = Pt(font_size_pt)
                        if show_answers and is_corr:
                            run_label.underline = True
                            run_label.font.color.rgb = RGBColor(185, 28, 28)
                        add_formatted_text(p_opt, content, font_name='Times New Roman', font_size_pt=font_size_pt)
                    else:
                        add_formatted_text(p_opt, str(opt), font_name='Times New Roman', font_size_pt=font_size_pt)

    # 5. Bảng đáp án tổng hợp (Chỉ in khi show_key_table = True)
    if show_key_table and questions:
        if is_compact:
            p_div_key = doc.add_paragraph()
            p_div_key.paragraph_format.space_before = Pt(12)
            p_div_key.paragraph_format.space_after = Pt(4)
            p_div_key.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_dk = p_div_key.add_run("=" * 45)
            run_dk.bold = True
            run_dk.font.color.rgb = RGBColor(100, 116, 139)
        else:
            doc.add_page_break()

        p_key_title = doc.add_paragraph()
        p_key_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_key_title.paragraph_format.space_before = Pt(4)
        p_key_title.paragraph_format.space_after = Pt(6)
        run_key = p_key_title.add_run("--- BẢNG ĐÁP ÁN ---")
        run_key.bold = True
        run_key.font.name = 'Times New Roman'
        run_key.font.size = Pt(12.5)
        run_key.font.color.rgb = RGBColor(22, 101, 52)

        cols = 10 if is_compact and len(questions) >= 10 else 5
        rows = (len(questions) + cols - 1) // cols
        table = doc.add_table(rows=rows * 2, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        for i, q in enumerate(questions):
            col_idx = i % cols
            row_idx = (i // cols) * 2

            cell_q = table.cell(row_idx, col_idx)
            cell_q.text = f"Câu {i+1}"
            p_cell_q = cell_q.paragraphs[0]
            p_cell_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if p_cell_q.runs:
                p_cell_q.runs[0].bold = True
                p_cell_q.runs[0].font.name = 'Times New Roman'
                p_cell_q.runs[0].font.size = Pt(9.5 if cols == 10 else 10)
            shading_q = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
            cell_q._tc.get_or_add_tcPr().append(shading_q)

            q_type = q.get('type', 'mcq')
            raw_ca = q.get('correct_answer')
            correct_display = ""

            if q_type == 'true_false' or isinstance(raw_ca, dict):
                if isinstance(raw_ca, dict):
                    parts = [f"{k}-{'Đ' if raw_ca.get(k) else 'S'}" for k in ['a', 'b', 'c', 'd'] if k in raw_ca]
                    correct_display = ", ".join(parts) if parts else str(raw_ca)
                else:
                    correct_display = str(raw_ca)
            elif q_type == 'short_answer':
                correct_display = str(raw_ca or '')
            else:
                correct_str = str(raw_ca or '').strip()
                for c in ["A", "B", "C", "D"]:
                    if correct_str.startswith(f"{c}.") or correct_str.startswith(f"{c}:") or correct_str == c:
                        correct_display = c
                        break
                if not correct_display and correct_str:
                    correct_display = correct_str[:2]

            cell_ans = table.cell(row_idx + 1, col_idx)
            cell_ans.text = correct_display or "-"
            p_cell_ans = cell_ans.paragraphs[0]
            p_cell_ans.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if p_cell_ans.runs:
                p_cell_ans.runs[0].bold = True
                p_cell_ans.runs[0].font.name = 'Times New Roman'
                p_cell_ans.runs[0].font.size = Pt(10 if cols == 10 else 11)
                p_cell_ans.runs[0].font.color.rgb = RGBColor(220, 38, 38)

    # 6. Hướng dẫn giải chi tiết (Chỉ in khi show_explains = True)
    has_explains = any(q.get('explain') for q in questions)
    if show_explains and has_explains:
        p_exp_div = doc.add_paragraph()
        p_exp_div.paragraph_format.space_before = Pt(10)
        p_exp_div.paragraph_format.space_after = Pt(4)
        run_exp_title = p_exp_div.add_run("HƯỚNG DẪN GIẢI CHI TIẾT:")
        run_exp_title.bold = True
        run_exp_title.font.name = 'Times New Roman'
        run_exp_title.font.size = Pt(11.5)
        run_exp_title.font.color.rgb = RGBColor(30, 41, 59)

        for idx, q in enumerate(questions, 1):
            exp = q.get('explain', '')
            if exp:
                p_exp = doc.add_paragraph()
                p_exp.paragraph_format.space_before = Pt(2 if is_compact else 4)
                p_exp.paragraph_format.space_after = Pt(2 if is_compact else 3)
                p_exp.paragraph_format.line_spacing = opt_line_spacing

                run_exp_q = p_exp.add_run(f"Câu {idx}: ")
                run_exp_q.bold = True
                run_exp_q.font.name = 'Times New Roman'
                run_exp_q.font.size = Pt(font_size_pt)
                run_exp_q.font.color.rgb = RGBColor(79, 70, 229)

                add_block_content(doc, exp, font_name='Times New Roman', font_size_pt=font_size_pt, first_p=p_exp, line_spacing=opt_line_spacing)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream
