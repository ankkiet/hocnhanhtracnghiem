import sys
import os
import unittest
from unittest.mock import MagicMock, patch

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

class TestFixedFeatures(unittest.TestCase):
    def test_teacher_analytics_with_dict_and_short_answer(self):
        """Kiểm tra get_quiz_analytics không bị crash bởi u_ans.strip() khi câu hỏi là Đúng/Sai hoặc Trả lời ngắn"""
        from routers.teacher import get_quiz_analytics

        mock_db = MagicMock()
        
        # Giả lập 2 bài nộp của học sinh
        sub1 = MagicMock()
        sub1.to_dict.return_value = {
            'student_name': 'Học sinh A',
            'score': 2.25,
            'total_questions': 3,
            'answers': {
                '0': 'A',  # MCQ
                '1': {'a': True, 'b': False, 'c': True, 'd': True},  # True/False (dict)
                '2': '3.14'  # Short answer
            }
        }
        
        sub2 = MagicMock()
        sub2.to_dict.return_value = {
            'student_name': 'Học sinh B',
            'score': 1.1,
            'total_questions': 3,
            'answers': {
                '0': 'b',  # MCQ lowercase
                '1': {'a': 'Đúng', 'b': 'Sai', 'c': False, 'd': False},  # True/False (mixed string/bool)
                '2': 3.14  # Float answer
            }
        }
        
        mock_quiz_doc = MagicMock()
        mock_quiz_doc.exists = True
        mock_quiz_doc.to_dict.return_value = {
            'title': 'Đề thi GDPT 2018',
            'data': [
                {
                    'question': 'Câu 1: MCQ',
                    'type': 'mcq',
                    'options': ['A. 1', 'B. 2', 'C. 3', 'D. 4'],
                    'correct_answer': 'A'
                },
                {
                    'question': 'Câu 2: Đúng Sai',
                    'type': 'true_false',
                    'options': ['a) Ý 1', 'b) Ý 2', 'c) Ý 3', 'd) Ý 4'],
                    'correct_answer': {'a': True, 'b': False, 'c': True, 'd': True}
                },
                {
                    'question': 'Câu 3: Điền số',
                    'type': 'short_answer',
                    'correct_answer': '3,14'
                }
            ]
        }
        
        mock_db.collection.return_value.document.return_value.get.return_value = mock_quiz_doc
        mock_db.collection.return_value.document.return_value.collection.return_value.get.return_value = [sub1, sub2]

        with patch('routers.teacher.get_db', return_value=mock_db), \
             patch('routers.teacher.verify_teacher_access', return_value=True):
            import asyncio
            result = asyncio.run(get_quiz_analytics(quiz_id='test_quiz_id', teacher_token='valid_token'))

        self.assertEqual(result['status'], 'success')
        self.assertEqual(result['total_submissions'], 2)
        # Verify no crash and accurate stats
        q_stats = result['question_stats']
        self.assertEqual(len(q_stats), 3)
        # Câu 1: sub1 correct, sub2 wrong -> 1 correct
        self.assertEqual(q_stats[0]['correct_count'], 1)
        # Câu 2: sub1 correct all 4 -> 1 correct
        self.assertEqual(q_stats[1]['correct_count'], 1)
        # Câu 3: both 3.14 and 3,14 matched -> 2 correct
        self.assertEqual(q_stats[2]['correct_count'], 2)

    def test_student_submit_exam_scoring_rules(self):
        """Kiểm tra quy chuẩn chấm điểm Bộ GD&ĐT: MCQ, Đúng/Sai (0.1, 0.25, 0.5, 1.0) và Trả lời ngắn"""
        from routers.student import submit_exam
        from pydantic import BaseModel

        class Req:
            quiz_id = 'test_quiz'
            student_name = 'Nguyễn Văn Test'
            student_token = 'token_abc'
            time_elapsed = 120
            answers = {
                '0': 'A',
                # 3 ý đúng (a, b, c) -> earned = 0.5
                '1': {'A': True, 'B': False, 'C': 'True', 'D': True},
                # Trả lời ngắn có dấu phẩy và khoảng trắng
                '2': '  12,5 '
            }

        mock_quiz_doc = MagicMock()
        mock_quiz_doc.exists = True
        mock_quiz_doc.to_dict.return_value = {
            'data': [
                {'type': 'mcq', 'correct_answer': 'A'},
                {'type': 'true_false', 'correct_answer': {'a': True, 'b': False, 'c': True, 'd': False}},
                {'type': 'short_answer', 'correct_answer': '12.5'}
            ]
        }

        mock_db = MagicMock()
        mock_db.collection.return_value.document.return_value.get.return_value = mock_quiz_doc
        mock_sub_ref = MagicMock()
        mock_db.collection.return_value.document.return_value.collection.return_value.document.return_value = mock_sub_ref

        with patch('routers.student.get_db', return_value=mock_db), \
             patch('routers.student.get_user_from_token', return_value={'id': 'user_123'}):
            import asyncio
            res = asyncio.run(submit_exam(Req()))

        self.assertEqual(res['status'], 'success')
        # Câu 1: 1.0, Câu 2: 0.5 (3 ý đúng a, b, c), Câu 3: 1.0 (12,5 khớp 12.5) => Total = 2.5
        self.assertEqual(res['score'], 2.5)
        self.assertEqual(res['results'][0]['earned'], 1.0)
        self.assertEqual(res['results'][1]['earned'], 0.5)
        self.assertEqual(res['results'][2]['earned'], 1.0)

    def test_student_progress_empty_id_guard(self):
        """Kiểm tra hàm get_progress và save_progress khi thiếu user_id không bị ném ValueError"""
        from routers.student import save_student_progress, get_student_progress

        class SaveReq:
            quiz_id = 'test_q'
            student_token = '   '  # whitespace
            progress_data = {'answers': {}}

        mock_db = MagicMock()
        with patch('routers.student.get_db', return_value=mock_db), \
             patch('routers.student.get_user_from_token', return_value=None):
            import asyncio
            save_res = asyncio.run(save_student_progress(SaveReq()))
            get_res = asyncio.run(get_student_progress('test_q', ''))

        self.assertEqual(save_res['status'], 'error')
        self.assertEqual(get_res['status'], 'success')
        self.assertIsNone(get_res['data'])

    def test_docx_export_gdpt2018_formats(self):
        """Kiểm tra export_quiz_to_docx định dạng chuẩn câu hỏi Đúng/Sai và Trả lời ngắn"""
        from core.docx_exporter import export_quiz_to_docx
        from docx import Document

        questions = [
            {
                "question": "Tính tích phân sau:",
                "type": "short_answer",
                "correct_answer": "42"
            },
            {
                "question": "Xét tính đúng sai của các khẳng định:",
                "type": "true_false",
                "options": ["a) Mệnh đề một", "b) Mệnh đề hai"],
                "correct_answer": {"a": True, "b": False}
            }
        ]

        stream = export_quiz_to_docx("ĐỀ THI TEST", questions, 45)
        doc = Document(stream)
        all_text = "\n".join(p.text for p in doc.paragraphs)

        # Kiểm tra sự xuất hiện của dòng điền đáp số cho câu trả lời ngắn
        self.assertIn("Đáp số:", all_text)
        # Kiểm tra sự xuất hiện của nhãn [ Đúng / Sai ] cho câu Đúng/Sai
        self.assertIn("[ Đúng / Sai ]", all_text)

    def test_auth_registration_role_sanitization(self):
        """Kiểm tra tài khoản thông thường không thể tự nâng cấp thành admin qua request"""
        from routers.auth import register
        
        class RegReq:
            username = 'regular_user@example.com'
            password = 'password123'
            full_name = 'Regular User'
            role = 'admin'  # Malicious attempt

        mock_db = MagicMock()
        mock_db.collection.return_value.where.return_value.get.return_value = []
        
        saved_doc = {}
        def mock_add(data):
            saved_doc.update(data)
            return (None, MagicMock())
        mock_db.collection.return_value.add = mock_add

        with patch('routers.auth.get_db', return_value=mock_db):
            import asyncio
            res = asyncio.run(register(RegReq()))

        self.assertEqual(res['status'], 'success')
        # Người dùng thường không thuộc ADMIN_EMAILS không thể nhận role admin
        self.assertEqual(saved_doc.get('role'), 'teacher')

    def test_answer_key_extraction_and_reconciliation(self):
        """Kiểm tra bóc tách Bảng đáp án chuẩn GDPT 2018 và áp đặt chính xác vào các loại câu hỏi"""
        from core.answer_key_extractor import (
            extract_answer_key_from_doc,
            separate_answer_key_from_text,
            reconcile_quiz_with_answer_key,
            AnswerKeyMap
        )

        # Giả lập văn bản đề thi có BẢNG ĐÁP ÁN ở cuối
        raw_exam_text = """
PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn.
Câu 1: Thủ đô của Việt Nam là gì?
A. Đà Nẵng
B. TP Hồ Chí Minh
C. Hà Nội
D. Hải Phòng

Câu 2: Số nguyên tố chẵn duy nhất là?
A. 0
B. 2
C. 4
D. 6

PHẦN II. Câu trắc nghiệm đúng sai.
Câu 1: Cho hàm số y = f(x).
a) Hàm số đồng biến trên R
b) Đồ thị hàm số đi qua gốc tọa độ
c) Hàm số có cực trị
d) Đạo hàm f'(x) luôn dương

PHẦN III. Câu trắc nghiệm trả lời ngắn.
Câu 1: Tìm giá trị lớn nhất của biểu thức P = -x^2 + 4x + 5.

--- HẾT ---

BẢNG ĐÁP ÁN
PHẦN I
1.C   2.B

PHẦN II
Câu 1: a-Đ, b-S, c-Đ, d-S

PHẦN III
Câu 1: 9
"""
        # 1. Tách phần Bảng đáp án ra khỏi đề thi
        q_text, ak_text, ak_map = separate_answer_key_from_text(raw_exam_text)
        self.assertNotIn("BẢNG ĐÁP ÁN", q_text)
        self.assertIn("BẢNG ĐÁP ÁN", ak_text)
        
        # 2. Kiểm tra các phần được trích xuất
        self.assertEqual(ak_map.part1.get(1), 'C')
        self.assertEqual(ak_map.part1.get(2), 'B')
        self.assertEqual(ak_map.part2.get(1), {'a': True, 'b': False, 'c': True, 'd': False})
        self.assertEqual(ak_map.part3.get(1), '9')

        # 3. Giả lập danh sách câu hỏi AI bóc tách (chưa có đáp án hoặc đáp án bị AI nhầm lẫn)
        unreconciled_quiz = [
            {
                "type": "mcq",
                "question": "Câu 1: Thủ đô của Việt Nam là gì?",
                "options": ["A. Đà Nẵng", "B. TP Hồ Chí Minh", "C. Hà Nội", "D. Hải Phòng"],
                "correct_answer": ""  # Chưa có đáp án
            },
            {
                "type": "mcq",
                "question": "Câu 2: Số nguyên tố chẵn duy nhất là?",
                "options": ["A. 0", "B. 2", "C. 4", "D. 6"],
                "correct_answer": "A. 0"  # AI bị ảo giác chọn A
            },
            {
                "type": "true_false",
                "question": "Câu 1: Cho hàm số y = f(x).",
                "options": ["a) Hàm số đồng biến trên R", "b) Đồ thị hàm số đi qua gốc tọa độ", "c) Hàm số có cực trị", "d) Đạo hàm f'(x) luôn dương"],
                "correct_answer": {"a": False, "b": False, "c": False, "d": False}
            },
            {
                "type": "short_answer",
                "question": "Câu 1: Tìm giá trị lớn nhất của biểu thức P = -x^2 + 4x + 5.",
                "options": [],
                "correct_answer": ""
            }
        ]

        # 4. Đối chiếu và áp đặt Bảng đáp án chính thức
        reconciled = reconcile_quiz_with_answer_key(unreconciled_quiz, ak_map)

        # Câu 1 (MCQ): Bảng đáp án là C -> Bắt buộc gán "C. Hà Nội"
        self.assertEqual(reconciled[0]["correct_answer"], "C. Hà Nội")
        # Câu 2 (MCQ): Bảng đáp án là B -> Bắt buộc sửa lại thành "B. 2" (ghi đè ảo giác của AI)
        self.assertEqual(reconciled[1]["correct_answer"], "B. 2")
        # Câu 1 Phần II (Đúng/Sai): Bắt buộc nhận {"a": True, "b": False, "c": True, "d": False}
        self.assertEqual(reconciled[2]["correct_answer"], {'a': True, 'b': False, 'c': True, 'd': False})
        # Câu 1 Phần III (Trả lời ngắn): Bắt buộc nhận '9'
        self.assertEqual(reconciled[3]["correct_answer"], '9')

    def test_extract_questions_from_text_bulletproof_with_bottom_answer_key(self):
        """Kiểm tra extract_questions_from_text_bulletproof không làm mất Bảng đáp án sau từ khóa HẾT và không gán bừa Option A"""
        from main import extract_questions_from_text_bulletproof

        text = """
Câu 1: Đâu là ngôn ngữ lập trình phổ biến cho AI?
A. HTML
B. CSS
C. Python
D. Photoshop

Câu 2: Số Pi có giá trị xấp xỉ bao nhiêu?
A. 2.17
B. 3.14
C. 1.41
D. 9.8

HẾT

BẢNG ĐÁP ÁN:
1.C 2.B
"""
        questions = extract_questions_from_text_bulletproof(text)
        self.assertEqual(len(questions), 2)
        # Kiểm tra đáp án được áp dụng từ BẢNG ĐÁP ÁN (1.C, 2.B), tuyệt đối không bị ép về Option A
        self.assertEqual(questions[0]["correct_answer"], "C. Python")
        self.assertEqual(questions[1]["correct_answer"], "B. 3.14")

    def test_rescue_true_false_from_short_answer_misclassification(self):
        """Kiểm tra hệ thống tự động phát hiện và phục hồi câu Đúng/Sai bị AI hoặc Parser gán nhầm thành Trả lời ngắn."""
        from services.ai_service import normalize_question_data
        from core.answer_key_extractor import reconcile_quiz_with_answer_key, AnswerKeyMap

        # Case 1: AI trả về type short_answer, options rỗng nhưng nội dung câu hỏi chứa 4 ý a, b, c, d
        q_raw1 = {
            "type": "short_answer",
            "question": "Câu 2: Cho hàm số f(x).\na) Hàm số đồng biến trên (0; 1).\nb) Đồ thị có 2 điểm cực trị.\nc) Giá trị nhỏ nhất là -4.\nd) Phương trình f(x) = 0 có 3 nghiệm.",
            "options": [],
            "correct_answer": "a. Đúng, b. Sai, c. Đúng, d. Sai"
        }
        res1 = normalize_question_data(q_raw1)
        self.assertEqual(res1["type"], "true_false")
        self.assertEqual(len(res1["options"]), 4)
        self.assertEqual(res1["options"][0], "a) Hàm số đồng biến trên (0; 1).")
        self.assertEqual(res1["correct_answer"], {"a": True, "b": False, "c": True, "d": False})

        # Case 2: AI trả về type short_answer nhưng correct_answer có dạng a-Đ, b-S, c-Đ, d-S
        q_raw2 = {
            "type": "short_answer",
            "question": "Cho tứ diện ABCD đều cạnh a.",
            "options": ["a) Góc giữa hai đường thẳng bằng 60 độ", "b) Thể tích bằng a^3/12", "c) Khoảng cách là a/2", "d) Bán kính mặt cầu ngoại tiếp là a*sqrt(6)/4"],
            "correct_answer": "a-Đ, b-S, c-Đ, d-S"
        }
        res2 = normalize_question_data(q_raw2)
        self.assertEqual(res2["type"], "true_false")
        self.assertEqual(len(res2["options"]), 4)
        self.assertEqual(res2["correct_answer"], {"a": True, "b": False, "c": True, "d": False})

        # Case 3: Thật sự là câu trả lời ngắn -> Giữ nguyên type short_answer
        q_raw3 = {
            "type": "short_answer",
            "question": "Tìm số nghiệm nguyên của bất phương trình.",
            "options": [],
            "correct_answer": "15,5"
        }
        res3 = normalize_question_data(q_raw3)
        self.assertEqual(res3["type"], "short_answer")
        self.assertEqual(res3["correct_answer"], "15.5")
        self.assertEqual(res3["options"], [])

        # Case 4: Reconcile áp dụng cứu hộ cho câu short_answer chứa các ý a-d
        ak = AnswerKeyMap()
        ak.set_answer(2, 1, {"a": True, "b": True, "c": False, "d": True})
        quiz = [{
            "type": "short_answer",
            "question": "Câu 1: Xét tính đúng sai của các mệnh đề:\na) 2 là số nguyên tố\nb) 4 là hợp số\nc) 1 là số nguyên tố\nd) 0 là số tự nhiên",
            "options": [],
            "correct_answer": ""
        }]
        reconciled = reconcile_quiz_with_answer_key(quiz, ak)
        self.assertEqual(reconciled[0]["type"], "true_false")
        self.assertEqual(len(reconciled[0]["options"]), 4)
        self.assertEqual(reconciled[0]["correct_answer"], {"a": True, "b": True, "c": False, "d": True})


    def test_azota_cach_3_underline_marking(self):
        """Kiểm tra chuẩn Azota Cách 3: Gạch chân trực tiếp đáp án đúng (<u>a)</u>, <u>a</u>), <u>a.</u>). Ký hiệu được gạch chân -> True, không gạch chân -> False."""
        from services.ai_service import clean_option_text, normalize_question_data

        # 1. Kiểm tra clean_option_text nhận diện chính xác gạch chân
        res_a = clean_option_text("<u>a)</u> Mệnh đề a đúng")
        self.assertEqual(res_a[3], "a")
        self.assertTrue(res_a[1])  # is_correct
        self.assertEqual(res_a[0], "Mệnh đề a đúng")

        res_b = clean_option_text("b) Mệnh đề b sai")
        self.assertEqual(res_b[3], "b")
        self.assertFalse(res_b[1])
        self.assertEqual(res_b[0], "Mệnh đề b sai")

        res_c = clean_option_text("<u>c</u>) Mệnh đề c đúng")
        self.assertEqual(res_c[3], "c")
        self.assertTrue(res_c[1])

        res_d = clean_option_text("<MARK><u>d.</u></MARK> Mệnh đề d đúng")
        self.assertEqual(res_d[3], "d")
        self.assertTrue(res_d[1])

        # 2. Kiểm tra normalize_question_data: Các ý không gạch chân BẮT BUỘC là False
        q_item = {
            "type": "true_false",
            "question": "Cho hàm số f(x).",
            "options": [
                "<u>a)</u> Hàm số đồng biến trên R",
                "b) Đồ thị có 2 điểm cực trị",
                "c) Giá trị lớn nhất bằng 5",
                "<u>d)</u> Đi qua điểm A(1; 2)"
            ]
        }
        norm = normalize_question_data(q_item)
        self.assertEqual(norm["type"], "true_false")
        self.assertEqual(norm["correct_answer"], {"a": True, "b": False, "c": False, "d": True})
        self.assertEqual(norm["options"][0], "a) Hàm số đồng biến trên R")
        self.assertEqual(norm["options"][1], "b) Đồ thị có 2 điểm cực trị")

    def test_azota_matrix_tags_preservation(self):
        """Kiểm tra tiền tố ma trận mức độ nhận thức Azota ([0, NB], [1, TH], [2, VD], [3, VDC]) không bị mất hay lặp."""
        from services.ai_service import clean_option_text, extract_sub_statements_from_text, normalize_question_data

        raw_text = """Cho hình chóp S.ABCD.
[0, NB] a) Đáy ABCD là hình vuông.
[1, TH] <u>b)</u> SA vuông góc với đáy.
[2, VD] c) Thể tích khối chóp bằng a^3/3.
[3, VDC] <u>d)</u> Khoảng cách từ A đến (SBD) bằng a*sqrt(2)/2."""

        stem, opts = extract_sub_statements_from_text(raw_text)
        self.assertEqual(stem, "Cho hình chóp S.ABCD.")
        self.assertEqual(len(opts), 4)
        self.assertTrue("[0, NB]" in opts[0])
        self.assertTrue("[1, TH]" in opts[1])

        # Chuẩn hóa qua normalize_question_data
        item = {
            "type": "true_false",
            "question": stem,
            "options": opts
        }
        norm = normalize_question_data(item)
        self.assertEqual(norm["options"][0], "a) [0, NB] Đáy ABCD là hình vuông.")
        self.assertEqual(norm["options"][1], "b) [1, TH] SA vuông góc với đáy.")
        self.assertEqual(norm["correct_answer"], {"a": False, "b": True, "c": False, "d": True})

    def test_azota_cach_2_per_question_tables(self):
        """Kiểm tra bóc tách bảng đáp án Đúng/Sai đặt ngay dưới câu hỏi chuẩn Azota Cách 2."""
        from docx import Document
        from core.docx_parser import parse_azota_tf_table

        doc = Document()
        # Bảng đứng: Lệnh hỏi, Đúng, Sai
        t_vert = doc.add_table(rows=5, cols=3)
        t_vert.rows[0].cells[0].text = "Lệnh hỏi"
        t_vert.rows[0].cells[1].text = "Đúng"
        t_vert.rows[0].cells[2].text = "Sai"

        t_vert.rows[1].cells[0].text = "a"
        t_vert.rows[1].cells[1].text = "x"
        t_vert.rows[1].cells[2].text = ""

        t_vert.rows[2].cells[0].text = "b"
        t_vert.rows[2].cells[1].text = ""
        t_vert.rows[2].cells[2].text = "x"

        t_vert.rows[3].cells[0].text = "c"
        t_vert.rows[3].cells[1].text = "v"
        t_vert.rows[3].cells[2].text = ""

        t_vert.rows[4].cells[0].text = "d"
        t_vert.rows[4].cells[1].text = ""
        t_vert.rows[4].cells[2].text = "1"

        res_vert = parse_azota_tf_table(t_vert)
        self.assertEqual(res_vert, {"a": True, "b": False, "c": True, "d": False})

        # Bảng ngang: Hàng 1: a | b | c | d, Hàng 2: Đ | S | Đ | S
        t_horiz = doc.add_table(rows=2, cols=4)
        t_horiz.rows[0].cells[0].text = "Ý a"
        t_horiz.rows[0].cells[1].text = "Ý b"
        t_horiz.rows[0].cells[2].text = "Ý c"
        t_horiz.rows[0].cells[3].text = "Ý d"

        t_horiz.rows[1].cells[0].text = "Đ"
        t_horiz.rows[1].cells[1].text = "S"
        t_horiz.rows[1].cells[2].text = "S"
        t_horiz.rows[1].cells[3].text = "Đ"

        res_horiz = parse_azota_tf_table(t_horiz)
        self.assertEqual(res_horiz, {"a": True, "b": False, "c": False, "d": True})


if __name__ == '__main__':
    unittest.main()
