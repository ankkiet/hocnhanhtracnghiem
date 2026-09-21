import os
import io
import unittest
import xml.etree.ElementTree as ET
from docx import Document
from fastapi.testclient import TestClient

import tempfile
from core.image_converter import detect_image_format, process_image_blob
from core.mathml_parser import parse_omath
from services.r2_service import upload_image_to_r2, get_stored_image
from main import app, extract_answer_key, extract_formatting_from_docx, parse_docx_to_marked_text


class TestBulletproofParsing(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.client = TestClient(app)

    def test_image_format_detection(self):
        """Test accurate sniffing of image magic bytes."""
        png_data = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR"
        jpeg_data = b"\xff\xd8\xff\xe0\x00\x10JFIF"
        gif_data = b"GIF89a"
        bmp_data = b"BM" + b"\x00" * 20
        emf_data = b"\x01\x00\x00\x00" + b"\x00" * 36 + b" EMF"
        wmf_data = b"\xd7\xcd\xc6\x9a" + b"\x00" * 20

        self.assertEqual(detect_image_format(png_data), "png")
        self.assertEqual(detect_image_format(jpeg_data), "jpeg")
        self.assertEqual(detect_image_format(gif_data), "gif")
        self.assertEqual(detect_image_format(bmp_data), "bmp")
        self.assertEqual(detect_image_format(emf_data), "emf")
        self.assertEqual(detect_image_format(wmf_data), "wmf")

    def test_image_blob_processing(self):
        """Test processing standard raster images without crashing."""
        tiny_png = (
            b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
            b"\x08\x06\x00\x00\x00\x1f\x15c4\x00\x00\x00\nIDATx\x9cc\x00\x01"
            b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
        )
        processed, ext = process_image_blob(tiny_png, "image/png")
        self.assertIn("png", ext.lower())
        self.assertTrue(len(processed) > 0)
        self.assertEqual(processed[:4], b"\x89PNG")

    def test_mathml_equation_arrays_and_fractions(self):
        """Test MathML parse_omath with fractions, superscripts and equation arrays."""
        from docx.oxml import parse_xml

        fraction_xml = """<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
            <m:f>
                <m:num><m:r><m:t>x</m:t></m:r></m:num>
                <m:den><m:r><m:t>2</m:t></m:r></m:den>
            </m:f>
        </m:oMath>"""
        elem = parse_xml(fraction_xml)
        latex = parse_omath(elem)
        self.assertIn(r"\frac{x}{2}", latex)

        eq_arr_xml = """<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
            <m:eqArr>
                <m:e><m:r><m:t>x = 1</m:t></m:r></m:e>
                <m:e><m:r><m:t>y = 2</m:t></m:r></m:e>
            </m:eqArr>
        </m:oMath>"""
        elem_arr = parse_xml(eq_arr_xml)
        latex_arr = parse_omath(elem_arr)
        self.assertIn(r"\begin{aligned}", latex_arr)
        self.assertIn(r"x = 1", latex_arr)
        self.assertIn(r"y = 2", latex_arr)
        self.assertIn(r"\end{aligned}", latex_arr)

    def test_extract_answer_key_from_docx_table(self):
        """Test extracting answer keys from bottom tables in DOCX."""
        doc = Document()
        doc.add_paragraph("Câu 1: Thủ đô của Việt Nam là gì?")
        doc.add_paragraph("A. Đà Nẵng")
        doc.add_paragraph("B. Hà Nội")
        doc.add_paragraph("C. TP.HCM")
        doc.add_paragraph("D. Cần Thơ")

        doc.add_paragraph("Câu 2: Số nguyên tố chẵn duy nhất là?")
        doc.add_paragraph("A. 0")
        doc.add_paragraph("B. 2")
        doc.add_paragraph("C. 4")
        doc.add_paragraph("D. 6")

        table = doc.add_table(rows=2, cols=2)
        row_q = table.rows[0].cells
        row_a = table.rows[1].cells
        row_q[0].text = "Câu 1"
        row_a[0].text = "B"
        row_q[1].text = "Câu 2"
        row_a[1].text = "B"

        full_text = "\n".join([p.text for p in doc.paragraphs])
        keys = extract_answer_key(doc, full_text)
        self.assertEqual(keys.get(1), "B")
        self.assertEqual(keys.get(2), "B")

    def test_docx_parsing_with_formatting_and_answer_table(self):
        """Test parsing questions from DOCX with formatting and linked answer table."""
        doc = Document()
        p1 = doc.add_paragraph()
        p1.add_run("Câu 1: Nước sôi ở bao nhiêu độ C?")
        doc.add_paragraph("A. 50")
        doc.add_paragraph("B. 80")
        doc.add_paragraph("C. 100")
        doc.add_paragraph("D. 120")

        # Answer key table
        table = doc.add_table(rows=2, cols=1)
        table.rows[0].cells[0].text = "Câu 1"
        table.rows[1].cells[0].text = "C"

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
            temp_path = tf.name
            doc.save(temp_path)

        try:
            questions = extract_formatting_from_docx(temp_path)
            self.assertTrue(len(questions) >= 1)
            self.assertIn("Nước sôi", questions[0]["question"])
            self.assertIn("C.", questions[0]["correct_answer"])
            self.assertEqual(len(questions[0]["options"]), 4)

            # Test text marking parser
            marked_text, img_map = parse_docx_to_marked_text(temp_path)
            self.assertIn("Nước sôi", marked_text)
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)

    def test_image_serving_endpoint(self):
        """Test caching and serving images via /api/images/{file_path:path}."""
        test_data = b"GIF89a\x01\x00\x01\x00\x80\x00\x00\xff\xff\xff\x00\x00\x00!\xf9\x04\x01\x00\x00\x00\x00,\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02D\x01\x00;"
        image_url = upload_image_to_r2(test_data, "test_img.gif")
        self.assertTrue(image_url.startswith("/api/images/"))

        response = self.client.get(image_url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, test_data)
        self.assertIn("image/gif", response.headers["content-type"])
        self.assertIn("max-age=31536000", response.headers.get("cache-control", ""))


    def test_wmf_emf_safe_handling(self):
        """Test that WMF/EMF binaries never raise unhandled exceptions and convert safely."""
        fake_emf = b"\x01\x00\x00\x00" + b"\x00" * 40
        data, mime = process_image_blob(fake_emf, "image/x-emf")
        self.assertIsNotNone(data)
        self.assertTrue(isinstance(data, bytes))

    def test_ai_audit_statistics_logic(self):
        """Test calculation of AI Quality Audit metrics and error breakdown."""
        total_questions = 10
        raw_issues = [
            {"question_index": 2, "category": "knowledge", "issue": "Sai công thức", "suggestion": "..."},
            {"question_index": 5, "category": "answer", "issue": "Đáp án C bị trùng", "suggestion": "..."},
            {"question_index": 8, "category": "grammar_typo", "issue": "Sai chính tả", "suggestion": "..."}
        ]
        
        category_counts = {
            "knowledge": 0,
            "answer": 0,
            "grammar_typo": 0,
            "format": 0
        }
        for item in raw_issues:
            cat = item.get("category", "knowledge")
            if cat in category_counts:
                category_counts[cat] += 1
            else:
                category_counts["knowledge"] += 1

        error_count = len(raw_issues)
        valid_questions = max(0, total_questions - error_count)
        accuracy_rate = round((valid_questions / total_questions) * 100, 1)

        self.assertEqual(error_count, 3)
        self.assertEqual(valid_questions, 7)
        self.assertEqual(accuracy_rate, 70.0)
        self.assertEqual(category_counts["knowledge"], 1)
        self.assertEqual(category_counts["answer"], 1)
        self.assertEqual(category_counts["grammar_typo"], 1)
        self.assertEqual(category_counts["format"], 0)


    def test_extract_questions_from_text_bulletproof(self):
        """Test bulletproof text extractor on unusual question and option formats."""
        from main import extract_questions_from_text_bulletproof

        text = """
        Câu 1 Cho hàm số y = f(x)
        A/ Giá trị 10
        B/ Giá trị 20
        C/ [ĐÚNG] Giá trị 30
        D/ Giá trị 40

        [Câu 2] Phương trình có nghiệm:
        (A) x = 1
        (B) x = 2
        (C) x = 3
        (D) x = 4
        """
        results = extract_questions_from_text_bulletproof(text)
        self.assertEqual(len(results), 2)
        self.assertEqual(len(results[0]["options"]), 4)
        self.assertIn("C.", results[0]["correct_answer"])
        self.assertEqual(len(results[1]["options"]), 4)

    def test_upload_docx_without_ai(self):
        """Test uploading a DOCX file without AI (use_ai=False) works 100% without error."""
        import time
        doc = Document()
        doc.add_paragraph("Câu 1: Mặt trời mọc ở hướng nào?")
        doc.add_paragraph("A. Đông")
        doc.add_paragraph("B. Tây")
        doc.add_paragraph("C. Nam")
        doc.add_paragraph("D. Bắc")

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        response = self.client.post(
            "/api/upload",
            files={"file": ("test_no_ai.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"use_ai": "false"}
        )
        self.assertEqual(response.status_code, 200)
        res_json = response.json()
        self.assertEqual(res_json["status"], "processing")
        task_id = res_json["task_id"]

        # Chờ task hoàn tất
        time.sleep(0.5)
        task_res = self.client.get(f"/api/task_status/{task_id}")
        self.assertEqual(task_res.status_code, 200)
        task_data = task_res.json()
        self.assertEqual(task_data["status"], "success")
        self.assertTrue(len(task_data["data"]) >= 1)
        self.assertIn("Mặt trời", task_data["data"][0]["question"])

    def test_upload_pdf_without_ai(self):
        """Test uploading a PDF file without AI (use_ai=False) parses locally instead of throwing 400 error."""
        import fitz
        import time

        # Tạo file PDF đơn giản bằng PyMuPDF
        pdf_doc = fitz.open()
        page = pdf_doc.new_page()
        page.insert_text(
            (50, 72),
            "Câu 1: 1 + 1 bằng bao nhiêu?\nA. 1\nB. 2\nC. 3\nD. 4\n\nCâu 2: 2 + 2 bằng bao nhiêu?\nA. 2\nB. 4\nC. 6\nD. 8"
        )
        pdf_bytes = pdf_doc.write()
        pdf_doc.close()

        response = self.client.post(
            "/api/upload",
            files={"file": ("test_no_ai.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
            data={"use_ai": "false"}
        )
        self.assertEqual(response.status_code, 200)
        res_json = response.json()
        self.assertEqual(res_json["status"], "processing")
        task_id = res_json["task_id"]

        time.sleep(0.5)
        task_res = self.client.get(f"/api/task_status/{task_id}")
        self.assertEqual(task_res.status_code, 200)
        task_data = task_res.json()
        self.assertEqual(task_data["status"], "success")
        self.assertEqual(len(task_data["data"]), 2)


if __name__ == "__main__":
    unittest.main()



