import os
import sys
import unittest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

import latex2mathml.converter
from docx import Document
from docx.oxml import parse_xml
from lxml import etree


class TestDocxMath(unittest.TestCase):
    def setUp(self):
        self.xsl_path = os.path.join(os.path.dirname(__file__), '..', 'core', 'MML2OMML.XSL')
        self.assertTrue(os.path.exists(self.xsl_path), "MML2OMML.XSL must exist")
        self.xslt = etree.XSLT(etree.parse(self.xsl_path))

    def test_omath_para(self):
        doc = Document()
        p = doc.add_paragraph('Test oMathPara: ')
        xml_str = '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:r><m:t>x = 1</m:t></m:r></m:oMath></m:oMathPara>'
        elem = parse_xml(xml_str.encode('utf-8'))
        p._p.append(elem)
        doc.save('test_omathpara.docx')
        self.assertTrue(os.path.exists('test_omathpara.docx'))
        os.remove('test_omathpara.docx')

    def test_latex_samples(self):
        samples = [
            r"f'(x)",
            r"y'",
            r"y''",
            r"(P): 2x - y + z - 1 = 0",
            r"d: \frac{x-1}{2} = \frac{y+2}{-1} = \frac{z}{3}",
            r"V = \frac{1}{3}Bh",
            r"\Delta ABC",
            r"\perp",
            r"\parallel",
            r"\angle A",
            r"\widehat{ABC}",
            r"\begin{aligned} x &= 1 \\ y &= 2 \end{aligned}",
            r"\begin{cases} x = 1 \\ y = 2 \end{cases}",
        ]
        parser = etree.XMLParser(recover=True)
        import re
        for s in samples:
            m = latex2mathml.converter.convert(s)
            m_fixed = re.sub(r'&(?!(?:amp|lt|gt|quot|apos|#x?[0-9a-fA-F]+);)', '&amp;', m)
            dom = etree.fromstring(m_fixed.encode('utf-8'), parser=parser)
            new_dom = self.xslt(dom)
            omml = etree.tostring(new_dom, encoding='utf-8')
            elem = parse_xml(omml)
            self.assertIsNotNone(elem)

    def test_export_quiz_to_docx(self):
        from core.docx_exporter import export_quiz_to_docx
        questions = [
            {
                "group_title": "PHẦN I: CÂU HỎI TRẮC NGHIỆM ĐẠI SỐ",
                "question": r"Cho hàm số \(y = \frac{x^2 - 3x + 2}{x - 1}\). Tìm tiệm cận xiên và giá trị của $m$ để $$x^2 - 4x + m = 0$$ có 2 nghiệm phân biệt.",
                "options": [
                    r"A. \(m < 4\)",
                    r"B. $m > 4$",
                    r"C. \(m = 4\)",
                    r"D. \(\forall m \in \mathbb{R}\)"
                ],
                "correct_answer": "A. m < 4",
                "explain": r"Biệt thức \(\Delta' = 4 - m\). Để phương trình có 2 nghiệm phân biệt thì \(\Delta' > 0 \Leftrightarrow 4 - m > 0 \Leftrightarrow m < 4\)."
            },
            {
                "question": r"Chất khí nào sinh ra khi cho Fe tác dụng với dung dịch HCl loãng: <b>Fe</b> + 2HCl \rightarrow FeCl<sub>2</sub> + H<sub>2</sub>?",
                "options": [
                    "A. Khí H2",
                    "B. Khí Cl2",
                    "C. Khí O2",
                    "D. Khí CO2"
                ],
                "correct_answer": "A",
                "explain": "Phản ứng giải phóng khí Hidro: Fe + 2HCl -> FeCl2 + H2."
            },
            {
                "question": r"Căn bậc hai của biểu thức \sqrt{x^2 + 2x + 1} bằng:",
                "options": [
                    r"A. \(|x + 1|\)",
                    r"B. \(x + 1\)",
                    r"C. \(-(x + 1)\)",
                    r"D. \((x + 1)^2\)"
                ],
                "correct_answer": "A",
                "explain": r"Vì \(\sqrt{(x+1)^2} = |x+1|\)."
            }
        ]
        stream = export_quiz_to_docx("ĐỀ THI KHẢO SÁT CHẤT LƯỢNG", questions, time_limit=45)
        self.assertIsNotNone(stream)
        stream_bytes = stream.getvalue()
        self.assertGreater(len(stream_bytes), 1000)

        # Kiểm tra đọc lại được tài liệu bằng docx.Document
        doc = Document(stream)
        self.assertGreaterEqual(len(doc.paragraphs), 5)
        self.assertGreaterEqual(len(doc.tables), 1)

        # Kiểm tra sự tồn tại của oMath trong XML
        full_xml = ""
        for p in doc.paragraphs:
            full_xml += p._p.xml
        self.assertIn("oMath", full_xml, "Văn bản Word phải chứa các đối tượng oMath (Word Equation)")
        # Đảm bảo không còn nguyên mã LaTeX thô như \frac hoặc \Delta trong văn bản thông thường
        self.assertNotIn(r"\frac", full_xml)
        self.assertNotIn(r"\Delta", full_xml)

    def test_export_quiz_to_docx_student_and_compact_options(self):
        from core.docx_exporter import export_quiz_to_docx

        questions = [
            {
                "question": "Giải phương trình 2x + 4 = 0",
                "options": ["A. x = -2", "B. x = 2", "C. x = 0", "D. x = 4"],
                "correct_answer": "A",
                "explain": "Ta có 2x = -4 suy ra x = -2."
            },
            {
                "question": "Tính tích phân sau:",
                "type": "short_answer",
                "correct_answer": "100"
            }
        ]

        # 1. Xuất bản cho học sinh (Student Mode)
        s_stream = export_quiz_to_docx(
            "ĐỀ THI HỌC SINH",
            questions,
            time_limit=15,
            target_audience="student",
            layout_mode="compact",
            include_student_answer_grid=True
        )
        self.assertIsNotNone(s_stream)
        s_doc = Document(s_stream)
        s_text = "\n".join(p.text for p in s_doc.paragraphs)
        self.assertIn("Họ và tên thí sinh:", s_text)
        self.assertIn("Lớp:", s_text)
        self.assertNotIn("--- BẢNG ĐÁP ÁN ---", s_text)
        self.assertNotIn("--- HƯỚNG DẪN GIẢI CHI TIẾT ---", s_text)
        self.assertNotIn("100", s_text)  # Không được lộ đáp án học sinh

        # 2. Xuất bản cho giáo viên (Teacher Mode)
        t_stream = export_quiz_to_docx(
            "ĐỀ THI GIÁO VIÊN",
            questions,
            time_limit=15,
            target_audience="teacher",
            layout_mode="standard",
            options_layout="4_col"
        )
        self.assertIsNotNone(t_stream)
        t_doc = Document(t_stream)
        t_text = "\n".join(p.text for p in t_doc.paragraphs)
        self.assertIn("--- BẢNG ĐÁP ÁN ---", t_text)
        self.assertIn("HƯỚNG DẪN GIẢI CHI TIẾT", t_text)
        self.assertIn("Ta có 2x = -4", t_text)

    def test_export_quiz_with_variation_table_and_html_table(self):
        """Kiểm tra xuất bảng biến thiên LaTeX chứa thẻ <br>, chuỗi \\nx và bảng HTML."""
        from core.docx_exporter import export_quiz_to_docx

        broken_table_q = (
            "Cho hàm số $y = f(x)$ có bảng biến thiên như sau:\\\n"
            "$$\\begin{array}{c|cccc} \\nx & -\\infty & -1 & 1 & +\\infty \\\\<br> "
            "y' & + & 0 & - & 0 & + \\\\<br> "
            "y & \\nearrow & 3 & \\searrow & -1 & \\nearrow <br>\\n "
            "\\end{array}$$Hàm số đạt cực đại tại điểm nào?"
        )

        questions = [
            {
                "question": broken_table_q,
                "options": ["A. x = -1", "B. x = 1", "C. y = 3", "D. y = -1"],
                "correct_answer": "A. x = -1",
                "explain": "Hàm số đạt cực đại tại x = -1 với giá trị y = 3."
            },
            {
                "question": "Bảng số liệu thống kê:<br><table><tr><th>Năm</th><th>2023</th><th>2024</th></tr><tr><td>Điểm</td><td>8.5</td><td>9.0</td></tr></table>Chọn nhận xét đúng:",
                "options": ["A. Điểm tăng", "B. Điểm giảm"],
                "correct_answer": "A. Điểm tăng"
            }
        ]

        stream = export_quiz_to_docx("ĐỀ THI TOÁN CÓ BẢNG BIẾN THIÊN", questions)
        self.assertIsNotNone(stream)
        doc = Document(stream)

        full_xml = "".join(p._p.xml for p in doc.paragraphs)
        for t in doc.tables:
            full_xml += t._tbl.xml

        # 1. Bắt buộc tạo được đối tượng công thức Office MathML (oMath)
        self.assertIn("oMath", full_xml)

        # 2. Không chứa rác thẻ <br> hay \nx trong văn bản
        import re
        text_nodes = re.findall(r'<(?:m:t|w:t)[^>]*>([^<]*)</(?:m:t|w:t)>', full_xml)
        all_text = " ".join(text_nodes)

        self.assertNotIn("<br>", all_text)
        self.assertNotIn("< br >", all_text)
        self.assertNotIn(r"\nx", all_text)
        self.assertNotIn("như sau:\\", all_text)
        self.assertIn("như sau:", all_text)

        # 3. Chứa bảng HTML được render chuẩn thành w:tbl trong Word
        self.assertGreaterEqual(len(doc.tables), 2)


if __name__ == '__main__':
    unittest.main()
