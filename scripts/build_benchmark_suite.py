"""
Script hoàn chỉnh sinh dữ liệu Benchmark:
- 30 file Word (.docx)
- 10 file PDF (.pdf)
- 40 file Ground Truth JSON (.json)
- 1 file manifest.json tổng hợp
"""

import json
import os
import re
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

import fitz
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from scripts.benchmark_cases_p1 import BENCHMARK_CASES
from scripts.benchmark_cases_p2 import BENCHMARK_CASES_P2

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BENCHMARK_DIR = os.path.join(BASE_DIR, "tests", "benchmark")
DATASET_DIR = os.path.join(BENCHMARK_DIR, "dataset")
DOCX_DIR = os.path.join(DATASET_DIR, "docx")
PDF_DIR = os.path.join(DATASET_DIR, "pdf")
GT_DIR = os.path.join(DATASET_DIR, "ground_truth")

os.makedirs(DOCX_DIR, exist_ok=True)
os.makedirs(PDF_DIR, exist_ok=True)
os.makedirs(GT_DIR, exist_ok=True)

FONT_PATH = "C:/Windows/Fonts/arial.ttf" if os.path.exists("C:/Windows/Fonts/arial.ttf") else None


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_docx_question(doc, q_item, style="underline"):
    q_num = q_item.get("question_number", 1)
    prefix = q_item.get("prefix", f"Câu {q_num}: ")
    q_text = q_item.get("question", "")
    q_type = q_item.get("type", "mcq")

    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_after = Pt(4)
    run_prefix = p_q.add_run(prefix)
    run_prefix.bold = True
    p_q.add_run(q_text)

    correct = q_item.get("correct_answer")

    if q_type == "mcq":
        options = q_item.get("options", [])
        if q_item.get("layout") == "single_line":
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.space_after = Pt(4)
            for idx, opt in enumerate(options):
                opt_char = opt[:1]
                is_correct = bool(correct and (correct.startswith(opt_char) or opt == correct))
                text_to_write = opt if not (style == "asterisk" and is_correct) else f"*{opt}"
                run = p_opt.add_run(text_to_write)
                if is_correct:
                    if style == "red":
                        run.font.color.rgb = RGBColor(255, 0, 0)
                    elif style == "bold":
                        run.bold = True
                    elif style not in ["asterisk", "bottom_table", "bottom_text"]:
                        run.underline = True
                if idx < len(options) - 1:
                    p_opt.add_run("\t\t")
        else:
            for opt in options:
                opt_char = opt[:1]
                is_correct = bool(correct and (correct.startswith(opt_char) or opt == correct))
                p_opt = doc.add_paragraph()
                p_opt.paragraph_format.space_after = Pt(3)
                
                text_to_write = opt if not (style == "asterisk" and is_correct) else f"*{opt}"
                
                m_opt = re.match(r'^(\*?[A-F][\.\:\)])\s*(.*)', text_to_write)
                if m_opt:
                    marker_str = m_opt.group(1)
                    content_str = m_opt.group(2)
                    r_marker = p_opt.add_run(marker_str + " ")
                    r_marker.bold = True
                    r_content = p_opt.add_run(content_str)
                    if is_correct:
                        if style == "red":
                            r_marker.font.color.rgb = RGBColor(255, 0, 0)
                            r_content.font.color.rgb = RGBColor(255, 0, 0)
                        elif style == "bold":
                            r_content.bold = True
                        elif style not in ["asterisk", "bottom_table", "bottom_text"]:
                            r_marker.underline = True
                            r_content.underline = True
                else:
                    r = p_opt.add_run(text_to_write)
                    if is_correct and style == "underline":
                        r.underline = True

    elif q_type == "true_false":
        options = q_item.get("options", [])
        tf_dict = correct if isinstance(correct, dict) else {}
        for opt in options:
            m_letter = re.match(r'^([a-d])[\)\.]\s*(.*)', opt, re.IGNORECASE)
            letter = m_letter.group(1).lower() if m_letter else "a"
            content = m_letter.group(2) if m_letter else opt
            
            p_opt = doc.add_paragraph()
            p_opt.paragraph_format.space_after = Pt(3)
            
            is_true = tf_dict.get(letter, False)
            marker_str = f"{letter}) "
            r_marker = p_opt.add_run(marker_str)
            r_marker.bold = True
            
            if style == "underline" and is_true:
                r_marker.underline = True
            elif style == "red" and is_true:
                r_marker.font.color.rgb = RGBColor(255, 0, 0)
            elif style == "bold" and is_true:
                r_marker.bold = True
                
            r_content = p_opt.add_run(content)
            if is_true and style not in ["underline", "red", "bold"]:
                r_marker.underline = True

    elif q_type == "short_answer":
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_after = Pt(4)
        r_ans = p_ans.add_run(f"Đáp án: {correct}")
        r_ans.italic = True

    if q_item.get("explain"):
        p_exp = doc.add_paragraph()
        p_exp.paragraph_format.space_after = Pt(6)
        r_lbl = p_exp.add_run("Lời giải chi tiết: ")
        r_lbl.bold = True
        r_lbl.italic = True
        p_exp.add_run(q_item["explain"])


def add_docx_answer_table(doc, answer_items):
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_title = doc.add_paragraph()
    r = p_title.add_run("BẢNG ĐÁP ÁN")
    r.bold = True
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    num_items = len(answer_items)
    table = doc.add_table(rows=2, cols=num_items)
    for idx, (q_num, ans) in enumerate(answer_items):
        cell_q = table.cell(0, idx)
        cell_q.text = f"Câu {q_num}"
        cell_q.paragraphs[0].runs[0].bold = True
        cell_q.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_margins(cell_q)

        cell_a = table.cell(1, idx)
        ans_clean = ans[:1] if ans and len(ans) >= 1 and ans[0].isupper() else str(ans)
        cell_a.text = ans_clean
        cell_a.paragraphs[0].runs[0].bold = True
        cell_a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_margins(cell_a)


def generate_pdf_file(filepath, title, questions, answer_text_block=None):
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    rect = fitz.Rect(45, 45, 550, 795)

    text_lines = []
    text_lines.append(title.upper())
    text_lines.append("=" * 45)
    text_lines.append("")

    for q in questions:
        q_num = q.get("question_number", 1)
        prefix = q.get("prefix", f"Câu {q_num}: ")
        text_lines.append(f"{prefix}{q.get('question', '')}")
        
        q_type = q.get("type", "mcq")
        if q_type == "mcq":
            for opt in q.get("options", []):
                text_lines.append(f"   {opt}")
        elif q_type == "true_false":
            for opt in q.get("options", []):
                text_lines.append(f"   {opt}")
        elif q_type == "short_answer":
            text_lines.append(f"   Đáp số: {q.get('correct_answer', '')}")
            
        if q.get("explain"):
            text_lines.append(f"   Lời giải: {q['explain']}")
        text_lines.append("")

    if answer_text_block:
        text_lines.append("BẢNG ĐÁP ÁN:")
        text_lines.append(answer_text_block)

    full_content = "\n".join(text_lines)

    if FONT_PATH:
        page.insert_font(fontname="arial", fontfile=FONT_PATH)
        page.insert_textbox(rect, full_content, fontname="arial", fontsize=10.5)
    else:
        page.insert_textbox(rect, full_content, fontsize=10.5)

    doc.save(filepath)
    doc.close()


def build_all():
    all_cases = BENCHMARK_CASES + BENCHMARK_CASES_P2
    print(f"[INFO] Bắt đầu tạo {len(all_cases)} bộ test cases...")

    manifest_entries = []
    subject_counts = {}
    format_counts = {"docx": 0, "pdf": 0}

    for case in all_cases:
        file_id = case["file_id"]
        file_name = case["file_name"]
        fmt = case["format_type"]
        subject = case["subject"]
        style = case["style"]
        questions = case["questions"]

        format_counts[fmt] = format_counts.get(fmt, 0) + 1
        subject_counts[subject] = subject_counts.get(subject, 0) + 1

        gt_data = {
            "file_id": file_id,
            "file_name": file_name,
            "format_type": fmt,
            "subject": subject,
            "style": style,
            "style_description": case.get("style_description", ""),
            "tags": case.get("tags", []),
            "total_questions": len(questions),
            "questions": questions
        }

        gt_path = os.path.join(GT_DIR, f"{file_id}.json")
        with open(gt_path, "w", encoding="utf-8") as f:
            json.dump(gt_data, f, ensure_ascii=False, indent=2)

        if fmt == "docx":
            out_path = os.path.join(DOCX_DIR, file_name)
            doc = docx.Document()

            p_title = doc.add_paragraph()
            r_title = p_title.add_run(f"ĐỀ THI KIỂM TRA MÔN {subject.upper()}")
            r_title.bold = True
            r_title.font.size = Pt(14)
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

            current_part = None
            last_group = None

            for q in questions:
                q_part = q.get("part")
                if q_part and q_part != current_part:
                    current_part = q_part
                    p_part = doc.add_paragraph()
                    p_part.paragraph_format.space_before = Pt(8)
                    p_part.paragraph_format.space_after = Pt(4)
                    if current_part == 1:
                        part_name = "PHẦN I: CÂU TRẮC NGHIỆM NHIỀU PHƯƠNG ÁN LỰA CHỌN"
                    elif current_part == 2:
                        part_name = "PHẦN II: CÂU TRẮC NGHIỆM ĐÚNG SAI"
                    else:
                        part_name = "PHẦN III: CÂU TRẮC NGHIỆM TRẢ LỜI NGẮN"
                    r_p = p_part.add_run(part_name)
                    r_p.bold = True

                group_title = q.get("group_title")
                if group_title and group_title != last_group:
                    last_group = group_title
                    p_grp = doc.add_paragraph()
                    p_grp.paragraph_format.space_before = Pt(6)
                    p_grp.paragraph_format.space_after = Pt(4)
                    r_grp = p_grp.add_run(group_title)
                    r_grp.italic = True

                add_docx_question(doc, q, style=style)

            if style == "bottom_table":
                answer_items = [(q["question_number"], q["correct_answer"]) for q in questions if q.get("type", "mcq") == "mcq"]
                if answer_items:
                    add_docx_answer_table(doc, answer_items)

            if case.get("answer_text_block"):
                doc.add_paragraph().paragraph_format.space_after = Pt(10)
                p_ak = doc.add_paragraph()
                r_h = p_ak.add_run("BẢNG ĐÁP ÁN:\n")
                r_h.bold = True
                p_ak.add_run(case["answer_text_block"])

            doc.save(out_path)

        elif fmt == "pdf":
            out_path = os.path.join(PDF_DIR, file_name)
            title = f"ĐỀ THI KIỂM TRA MÔN {subject}"
            ans_block = case.get("answer_text_block")
            generate_pdf_file(out_path, title, questions, answer_text_block=ans_block)

        manifest_entries.append({
            "file_id": file_id,
            "file_name": file_name,
            "format_type": fmt,
            "subject": subject,
            "style": style,
            "style_description": case.get("style_description", ""),
            "tags": case.get("tags", []),
            "total_questions": len(questions)
        })

    manifest_data = {
        "dataset_name": "HocNhanhTN Parser Benchmark Dataset",
        "version": "1.0.0",
        "total_files": len(all_cases),
        "format_counts": format_counts,
        "subject_counts": subject_counts,
        "total_questions": sum(e["total_questions"] for e in manifest_entries),
        "files": manifest_entries
    }

    manifest_path = os.path.join(BENCHMARK_DIR, "manifest.json")
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest_data, f, ensure_ascii=False, indent=2)

    print(f"[SUCCESS] Đã tạo thành công {len(all_cases)} files:")
    print(f" - Word (.docx): {format_counts['docx']} files")
    print(f" - PDF (.pdf): {format_counts['pdf']} files")
    print(f" - Ground Truth (.json): {len(all_cases)} files")
    print(f" - Tổng số câu hỏi trong dataset: {manifest_data['total_questions']} câu")
    print(f" - Manifest ghi tại: {manifest_path}")


if __name__ == "__main__":
    build_all()
